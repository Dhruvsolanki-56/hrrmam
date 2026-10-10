"""Excel exports (built from the same dicts the screens use, so filters match)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import workflow

STATUS = {"active": "In progress", "on_hold": "On hold", "rejected": "Rejected", "completed": "Completed"}
RED = Font(bold=True, color="A1260E")


def _sheet(title: str, head: list[str], rows: list[list], widths: list[int], flag_col: int | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(head)
    for r in rows:
        ws.append(r)
    fill = PatternFill("solid", fgColor="004248")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if flag_col is not None:
        for row in ws.iter_rows(min_row=2):
            if row[flag_col].value == "Yes":
                row[flag_col].font = RED
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def projects_xlsx(rows: list[dict]) -> bytes:
    head = ["Code", "Project", "Product", "Market", "Type", "Project manager", "Submission type", "Stage #", "Stage",
            "Also active (parallel)", "Status", "Health", "Critical blocker", "Open tasks", "Overdue tasks", "Overdue stage",
            "Target submission", "Target launch", "Last updated"]
    out = []
    for r in rows:
        s = r["stage"]
        out.append([r["code"] or "Pending", r["name"], r["product"], r["market"], r["project_type"],
                    (r["project_manager"] or {}).get("name", ""), r["submission_type"], s["phase"], s["name"],
                    ", ".join(r["stage_names"][k] for k in r["active_stages"] if k != r["stage_key"]),
                    STATUS[r["status"]], r["health"].title(), r["critical_blocker"], r["open_tasks"], r["overdue_tasks"],
                    "Yes" if r["overdue"] else "No", r["target_submission"] or "", r["target_launch"] or "", (r["updated_at"] or "")[:10]])
    return _sheet("Projects", head, out, [14, 38, 24, 12, 13, 18, 34, 8, 26, 30, 12, 9, 30, 10, 10, 10, 14, 14, 13], 15)


def legal_xlsx(rows: list[dict]) -> bytes:
    head = ["Project code", "Project", "Agreement", "Counterparty", "Status", "Owner", "Due", "Signed", "Age (days, unsigned)", "Overdue", "Link"]
    out = [[r["project_code"] or "", r["project_name"], r["label"], r["counterparty"], r["status"].replace("_", " ").title(),
            (r["owner"] or {}).get("name", ""), r["due_date"] or "", r["signed_date"] or "", r["age_days"] if r["age_days"] is not None else "",
            "Yes" if r["overdue"] else "No", r["link"]] for r in rows]
    return _sheet("Legal register", head, out, [14, 38, 20, 26, 18, 18, 12, 12, 12, 9, 40], 9)
